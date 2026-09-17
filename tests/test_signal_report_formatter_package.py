"""Regression tests for Issue #285 Stage 2, step 1: moving
`format_signal_report`'s implementation from the root-level
`report_formatter.py` into `ai_asset_platform.reports.signal_report_formatter`,
with `report_formatter.py` kept as a backward-compatibility shim and
`signal_runner.py` updated to import the package implementation directly.

Behavior (arguments, return value, exceptions, Excel output) is unchanged;
only the import location moved.
"""
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT_DIR = Path(__file__).resolve().parents[1]


def test_format_signal_report_importable_from_package():
    """The implementation lives in the package now."""
    from ai_asset_platform.reports.signal_report_formatter import (
        format_signal_report,
    )

    assert callable(format_signal_report)


def test_root_report_formatter_is_a_compatibility_shim():
    """`report_formatter.format_signal_report` still works, and is the
    *same* function object as the package implementation -- not a
    reimplementation that could drift out of sync.
    """
    import report_formatter
    from ai_asset_platform.reports.signal_report_formatter import (
        format_signal_report as package_format_signal_report,
    )

    assert report_formatter.format_signal_report is package_format_signal_report


def _write_sample_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Ticker",
            "Signal",
            "Rank",
            "Score",
            "Close",
            "Reason",
        ]
    )
    sheet.append(["AAPL", "BUY", 1, 95, 123.45, "テクニカル判定とAI判定が一致"])
    sheet.append(["SPY", "SELL", 2, 80, 456.7, "AI最終判定"])
    sheet.append(["9432.T", "HOLD", 5, 10, 166.1, "判定不一致のため安全側でHOLD"])
    workbook.save(path)


def test_format_signal_report_real_xlsx_round_trip(tmp_path):
    """Build a real temporary .xlsx, run `format_signal_report` on it (via
    the package import), and confirm the documented formatting contract:
    sheet renamed, header frozen, BUY/SELL/HOLD rows all handled without
    exception, and the file is still openable afterward.
    """
    from ai_asset_platform.reports.signal_report_formatter import (
        format_signal_report,
    )

    output_path = tmp_path / "signal_report.xlsx"
    _write_sample_workbook(output_path)

    format_signal_report(output_path)

    reopened = openpyxl.load_workbook(output_path)
    sheet = reopened.active
    assert sheet.title == "最新シグナル"
    assert sheet.freeze_panes == "A2"
    # 3 sample rows (BUY/SELL/HOLD) + header, all survived the round trip.
    assert sheet.max_row == 4


def test_format_signal_report_accepts_str_path(tmp_path):
    """The public signature accepts `str`, not just `Path` -- unchanged
    from the original root-level implementation.
    """
    from ai_asset_platform.reports.signal_report_formatter import (
        format_signal_report,
    )

    output_path = tmp_path / "signal_report_str.xlsx"
    _write_sample_workbook(output_path)

    format_signal_report(str(output_path))

    reopened = openpyxl.load_workbook(output_path)
    assert reopened.active.title == "最新シグナル"


def test_package_import_does_not_require_repo_root_on_sys_path(tmp_path):
    """Isolated regression, offline (no network, no venv creation, no
    dependency reinstall): import
    `ai_asset_platform.reports.signal_report_formatter` in a subprocess
    whose cwd is *not* the repo root and whose PYTHONPATH is set to
    `src/` only -- never the repo root itself, so root-level modules like
    `report_formatter.py` or `signal_runner.py` are never on the import
    path at all. This proves resolution *correctness* (it resolves to
    this checkout's own package file, not the root-level shim, regardless
    of cwd), not resolution *mechanism*.

    Deliberately does NOT assert that removing PYTHONPATH breaks the
    import. In CI, `pip install -e .` already ran before this test suite,
    so `ai_asset_platform` resolves via the editable install regardless
    of PYTHONPATH -- asserting the opposite would be an environment-
    dependent false failure there, not a real regression in this
    checkout. (A prior version of this test asserted exactly that and
    failed in CI for this reason.)
    """
    src_dir = ROOT_DIR / "src"
    expected_module_path = (
        src_dir / "ai_asset_platform" / "reports" / "signal_report_formatter.py"
    ).resolve()
    assert expected_module_path.is_file()

    env = dict(os.environ)
    # Only ever add src/ -- never the repo root -- so root-level modules
    # stay unreachable regardless of what else this interpreter already
    # has installed (e.g. an unrelated checkout's editable install, if
    # this test happens to run under a borrowed/shared venv).
    env["PYTHONPATH"] = str(src_dir)

    script = (
        "import ai_asset_platform.reports.signal_report_formatter as m\n"
        "import sys\n"
        "assert 'report_formatter' not in sys.modules, ("
        "'root-level report_formatter must not have been imported')\n"
        "print(m.__file__)\n"
    )
    result = subprocess.run(
        [sys.executable, "-c", script],
        cwd=str(tmp_path),
        env=env,
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert result.returncode == 0, (
        f"isolated package import failed:\nstdout={result.stdout}\nstderr={result.stderr}"
    )

    resolved_module_path = Path(result.stdout.strip()).resolve()
    assert resolved_module_path == expected_module_path, (
        "ai_asset_platform.reports.signal_report_formatter resolved to a "
        f"different file than this checkout's own module: {resolved_module_path} "
        f"(expected {expected_module_path})"
    )


def test_signal_runner_does_not_import_from_root_report_formatter():
    """Static regression: `signal_runner.py` must import the package
    implementation directly (`ai_asset_platform.reports.signal_report_formatter`),
    not the root-level compatibility shim (`from report_formatter import ...`).
    """
    signal_runner_path = ROOT_DIR / "signal_runner.py"
    text = signal_runner_path.read_text(encoding="utf-8")

    assert "from report_formatter import" not in text, (
        "signal_runner.py must not import from the root report_formatter shim"
    )
    assert "from ai_asset_platform.reports.signal_report_formatter import" in text, (
        "signal_runner.py must import format_signal_report from the package "
        "implementation directly"
    )


def test_format_signal_report_missing_file_raises(tmp_path):
    """Exception behavior for a nonexistent input file is unchanged
    (openpyxl's own FileNotFoundError propagates, not swallowed).
    """
    from ai_asset_platform.reports.signal_report_formatter import (
        format_signal_report,
    )

    missing_path = tmp_path / "does_not_exist.xlsx"
    with pytest.raises(FileNotFoundError):
        format_signal_report(missing_path)
